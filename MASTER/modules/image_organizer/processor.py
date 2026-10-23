import os
import re
import shutil
import datetime

from core.image_manager import ImageManager
from core.filename_builder import FilenameBuilder


def extract_image_path(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    val = cell.value
    if isinstance(val, str):
        m = re.search(r'HYPERLINK\(\s*["\']([^"\']+)', val, re.I)
        if m:
            return m.group(1)
        m = re.search(r'(\\\\[^"\s]+\.(jpg|jpeg|png|tif|bmp|webp))', val, re.I)
        if m:
            return m.group(1)
        m = re.search(r'([A-Za-z]:\\[^"\s]+\.(jpg|jpeg|png|tif|bmp|webp))', val, re.I)
        if m:
            return m.group(1)
    return None


def normalize_path(p):
    if not p:
        return None
    p = p.strip().replace("/", "\\")
    p = re.sub(r'^file:/*', '', p, flags=re.I)
    return p


def clean_cell(val):
    if val is None:
        return None
    s = str(val).strip()
    if s.startswith('='):
        return None
    return s if s else None


def get_col_index(headers, header_name):
    """Return 1-based column index for a header name, or None."""
    if not header_name:
        return None
    for i, h in enumerate(headers):
        if h and str(h).strip().upper() == header_name.strip().upper():
            return i + 1  # openpyxl is 1-based
    return None


class Processor:

    def __init__(self, project, logger=None, progress_callback=None, done_callback=None):
        self.project = project
        self.logger = logger
        self.progress_callback = progress_callback
        self.done_callback = done_callback
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def log(self, text):
        if self.logger:
            self.logger.write(text)
        print(text)

    def set_progress(self, value):
        if self.progress_callback:
            self.progress_callback(value)

    # --------------------------------------------------
    # MAIN ENTRY
    # --------------------------------------------------

    def run(self):
        import openpyxl

        p = self.project
        output = p.output_folder
        if not output:
            self.log("ERROR: No output folder selected.")
            return

        os.makedirs(output, exist_ok=True)

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        log_path = os.path.join(output, f"log_{timestamp}.txt")

        self.log("=" * 50)
        self.log(f"IMAGE ORGANIZER PRO  - MODE {p.mode}")
        self.log(f"Output : {output}")
        self.log(f"Files  : {len(p.excel_files)}")
        self.log("=" * 50)

        total_copied = 0
        total_skipped = 0

        all_rows = self._count_rows(p.excel_files)
        processed = 0

        for excel_file in p.excel_files:
            if self._cancelled:
                break
            if not os.path.exists(excel_file):
                self.log(f"NOT FOUND: {excel_file}")
                continue

            self.log(f"\nFILE: {excel_file}")
            excel_dir = os.path.dirname(excel_file)

            wb = openpyxl.load_workbook(excel_file, data_only=False)

            for ws in wb.worksheets:
                if self._cancelled:
                    break
                self.log(f"  Sheet: {ws.title}")

                if ws.max_row < 2:
                    continue

                # Read headers from row 1
                headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

                for row in range(2, ws.max_row + 1):
                    if self._cancelled:
                        break
                    try:
                        ok = self._process_row(ws, row, headers, excel_dir, p)
                        if ok:
                            total_copied += 1
                        else:
                            total_skipped += 1
                    except Exception as e:
                        self.log(f"  Row {row} ERROR: {e}")
                        total_skipped += 1

                    processed += 1
                    if all_rows > 0:
                        self.set_progress(int(processed / all_rows * 100))

            wb.close()

        self.set_progress(100)
        self.log("\n" + "=" * 50)
        self.log("DONE")
        self.log(f"COPIED : {total_copied}")
        self.log(f"SKIPPED: {total_skipped}")
        self.log("=" * 50)

        p.copied_images = total_copied
        p.missing_images = total_skipped

        if self.done_callback:
            self.done_callback(total_copied, total_skipped)

    # --------------------------------------------------
    # ROW DISPATCH
    # --------------------------------------------------

    def _process_row(self, ws, row, headers, excel_dir, p):
        mode = p.mode
        m = p.header_mapping

        # --- Resolve image path ---
        img_col = get_col_index(headers, m.get("IMAGE_PATH"))
        if not img_col:
            return False

        cell = ws.cell(row=row, column=img_col)
        raw = extract_image_path(cell)
        if not raw:
            raw = clean_cell(cell.value)
        img = normalize_path(raw)
        if not img:
            return False

        if not os.path.isabs(img):
            img = os.path.join(excel_dir, img)

        if not os.path.exists(img):
            self.log(f"  MISSING row {row}: {img}")
            return False

        ext = os.path.splitext(img)[1]

        # --- Build filename from mapped fields ---
        row_data = {}
        for field in p.filename_fields:
            col = get_col_index(headers, m.get(field, field))
            if col:
                row_data[field] = clean_cell(ws.cell(row=row, column=col).value) or ""
            else:
                row_data[field] = ""

        filename = FilenameBuilder.build_filename(row_data, p.filename_fields, ext)

        # --- Determine destination folder ---
        if mode == 1:
            dest = self._mode1_folder(ws, row, headers, p)
        elif mode == 2:
            dest = self._mode2_folder(ws, row, headers, p)
        elif mode == 3:
            dest = self._mode3_folder(ws, row, headers, p)
        else:
            dest = p.output_folder

        ImageManager.copy_image(img, dest, filename)
        self.log(f"  OK row {row}: {filename}")
        return True

    # --------------------------------------------------
    # MODE 1: NESTED FOLDERS from column range
    # --------------------------------------------------

    def _mode1_folder(self, ws, row, headers, p):
        m = p.header_mapping
        folder = p.output_folder
        last = None

        folder_cols = m.get("FOLDER_COLS", [])  # list of header names
        for header_name in folder_cols:
            col = get_col_index(headers, header_name)
            if not col:
                continue
            val = clean_cell(ws.cell(row=row, column=col).value)
            if val:
                safe = re.sub(r'[<>:"/\\|?*]', '_', val)
                folder = os.path.join(folder, safe)
                last = folder

        return last if last else os.path.join(p.output_folder, "UNCATEGORIZED")

    # --------------------------------------------------
    # MODE 2: KEEP/DELETE + COLLECTION
    # --------------------------------------------------

    def _mode2_folder(self, ws, row, headers, p):
        m = p.header_mapping

        kd_col = get_col_index(headers, m.get("KD"))
        kd_val = ""
        if kd_col:
            kd_val = clean_cell(ws.cell(row=row, column=kd_col).value) or ""

        first = "DELETE" if "DELETE" in kd_val.upper() else "KEEP"

        coll_col = get_col_index(headers, m.get("COLLECTION NAME"))
        coll_val = ""
        if coll_col:
            coll_val = clean_cell(ws.cell(row=row, column=coll_col).value) or ""

        second = re.sub(r'[<>:"/\\|?*]', '_', coll_val) if coll_val else "SOLE"

        return os.path.join(p.output_folder, first, second)

    # --------------------------------------------------
    # MODE 3: KEEP/DELETE + MULTI-LEVEL CUSTOM
    # --------------------------------------------------

    def _mode3_folder(self, ws, row, headers, p):
        m = p.header_mapping

        kd_col = get_col_index(headers, m.get("KD"))
        kd_val = ""
        if kd_col:
            kd_val = clean_cell(ws.cell(row=row, column=kd_col).value) or ""

        top = "DELETE" if "DELETE" in kd_val.upper() else "KEEP"

        folder_cols = m.get("FOLDER_COLS", [])
        parts = [top]
        for header_name in folder_cols:
            col = get_col_index(headers, header_name)
            if not col:
                continue
            val = clean_cell(ws.cell(row=row, column=col).value)
            if val:
                parts.append(re.sub(r'[<>:"/\\|?*]', '_', val))

        return os.path.join(p.output_folder, *parts)

    # --------------------------------------------------
    # HELPER
    # --------------------------------------------------

    def _count_rows(self, files):
        import openpyxl
        total = 0
        for f in files:
            if not os.path.exists(f):
                continue
            try:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    total += max(0, ws.max_row - 1)
                wb.close()
            except Exception:
                pass
        return total
