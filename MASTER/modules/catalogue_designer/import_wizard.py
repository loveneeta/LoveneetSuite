# ==========================================================
# CATALOGUE DESIGNER — Import Wizard
# Step 1: Select image folder and Excel file
# ==========================================================

import os
import sys
import threading
from pathlib import Path
from typing import Optional, Callable

import customtkinter as ctk
from tkinter import filedialog, messagebox

# Add parent for imports
_HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _HERE)

from data_models import ImageFile, ExcelMetadata, DesignSession
import openpyxl


# ──────────────────────────────────────────────────────────
# CONSTANTS
# ──────────────────────────────────────────────────────────

VALID_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
VALID_EXCEL_EXTENSIONS = {".xlsx", ".xls"}


# ──────────────────────────────────────────────────────────
# IMPORT WIZARD
# ──────────────────────────────────────────────────────────

class ImportWizard(ctk.CTkFrame):
    """Step 1: Import images and Excel file"""
    
    def __init__(self, parent, on_complete: Callable = None):
        super().__init__(parent)
        
        self.on_complete = on_complete  # Callback when import succeeds
        
        # ─────────────────────────────────────
        # STATE
        # ─────────────────────────────────────
        
        self.image_folder = ""
        self.excel_file = ""
        self.images: list[ImageFile] = []
        self.excel_metadata: Optional[ExcelMetadata] = None
        
        self.is_loading = False
        
        # ─────────────────────────────────────
        # BUILD UI
        # ─────────────────────────────────────
        
        self.build_ui()
    
    # ══════════════════════════════════════════════════════
    # UI CONSTRUCTION
    # ══════════════════════════════════════════════════════
    
    def build_ui(self):
        """Build the import wizard UI"""
        
        # ─────────────────────────────────────
        # TITLE
        # ─────────────────────────────────────
        
        title = ctk.CTkLabel(
            self,
            text="Step 1: Import Images & Excel",
            font=("Arial", 24, "bold")
        )
        
        title.pack(pady=(20, 10), padx=20)
        
        desc = ctk.CTkLabel(
            self,
            text="Select an image folder and Excel file to begin",
            font=("Arial", 14),
            text_color="#AAAAAA"
        )
        
        desc.pack(pady=(0, 30))
        
        # ─────────────────────────────────────
        # MAIN CONTAINER
        # ─────────────────────────────────────
        
        main_frame = ctk.CTkFrame(self)
        main_frame.pack(fill="both", expand=True, padx=30, pady=20)
        
        # ─────────────────────────────────────
        # IMAGE FOLDER SECTION
        # ─────────────────────────────────────
        
        img_frame = ctk.CTkFrame(main_frame)
        img_frame.pack(fill="x", pady=(0, 25))
        
        img_label = ctk.CTkLabel(
            img_frame,
            text="📁 Image Folder",
            font=("Arial", 16, "bold")
        )
        img_label.pack(anchor="w", pady=(0, 10))
        
        img_path_frame = ctk.CTkFrame(img_frame)
        img_path_frame.pack(fill="x", pady=(0, 10))
        
        self.img_path_display = ctk.CTkLabel(
            img_path_frame,
            text="No folder selected",
            font=("Arial", 12),
            text_color="#888888"
        )
        self.img_path_display.pack(anchor="w", side="left", expand=True)
        
        img_btn_frame = ctk.CTkFrame(img_frame)
        img_btn_frame.pack(fill="x")
        
        browse_img_btn = ctk.CTkButton(
            img_btn_frame,
            text="Browse Folder",
            command=self.select_image_folder,
            width=120,
            height=35
        )
        browse_img_btn.pack(side="left", padx=(0, 10))
        
        self.img_count_label = ctk.CTkLabel(
            img_btn_frame,
            text="",
            font=("Arial", 11),
            text_color="#AAAAAA"
        )
        self.img_count_label.pack(side="left")
        
        # ─────────────────────────────────────
        # EXCEL FILE SECTION
        # ─────────────────────────────────────
        
        excel_frame = ctk.CTkFrame(main_frame)
        excel_frame.pack(fill="x", pady=(0, 25))
        
        excel_label = ctk.CTkLabel(
            excel_frame,
            text="📊 Excel File",
            font=("Arial", 16, "bold")
        )
        excel_label.pack(anchor="w", pady=(0, 10))
        
        excel_path_frame = ctk.CTkFrame(excel_frame)
        excel_path_frame.pack(fill="x", pady=(0, 10))
        
        self.excel_path_display = ctk.CTkLabel(
            excel_path_frame,
            text="No file selected",
            font=("Arial", 12),
            text_color="#888888"
        )
        self.excel_path_display.pack(anchor="w", side="left", expand=True)
        
        excel_btn_frame = ctk.CTkFrame(excel_frame)
        excel_btn_frame.pack(fill="x")
        
        browse_excel_btn = ctk.CTkButton(
            excel_btn_frame,
            text="Browse File",
            command=self.select_excel_file,
            width=120,
            height=35
        )
        browse_excel_btn.pack(side="left", padx=(0, 10))
        
        self.excel_info_label = ctk.CTkLabel(
            excel_btn_frame,
            text="",
            font=("Arial", 11),
            text_color="#AAAAAA"
        )
        self.excel_info_label.pack(side="left")
        
        # ─────────────────────────────────────
        # STATUS / LOADING
        # ─────────────────────────────────────
        
        self.status_label = ctk.CTkLabel(
            main_frame,
            text="",
            font=("Arial", 11),
            text_color="#AAAAAA"
        )
        self.status_label.pack(anchor="w", pady=(20, 0))
        
        # ─────────────────────────────────────
        # ACTION BUTTONS
        # ─────────────────────────────────────
        
        button_frame = ctk.CTkFrame(self)
        button_frame.pack(fill="x", padx=30, pady=(0, 30))
        
        cancel_btn = ctk.CTkButton(
            button_frame,
            text="Cancel",
            command=self.on_cancel,
            width=120,
            height=40,
            fg_color="#333333"
        )
        cancel_btn.pack(side="right", padx=(10, 0))
        
        self.next_btn = ctk.CTkButton(
            button_frame,
            text="Next →",
            command=self.on_next,
            width=120,
            height=40,
            state="disabled"
        )
        self.next_btn.pack(side="right", padx=(10, 0))
    
    # ══════════════════════════════════════════════════════
    # FILE SELECTION
    # ══════════════════════════════════════════════════════
    
    def select_image_folder(self):
        """Browse and select image folder"""
        
        folder = filedialog.askdirectory(
            title="Select Image Folder",
            initialdir=os.path.expanduser("~")
        )
        
        if folder:
            self.image_folder = folder
            self.update_ui()
    
    def select_excel_file(self):
        """Browse and select Excel file"""
        
        file = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel Files", "*.xlsx *.xls"), ("All Files", "*.*")],
            initialdir=os.path.expanduser("~")
        )
        
        if file:
            self.excel_file = file
            self.update_ui()
    
    # ══════════════════════════════════════════════════════
    # UPDATE UI
    # ══════════════════════════════════════════════════════
    
    def update_ui(self):
        """Update UI based on current selections"""
        
        # Image folder display
        if self.image_folder:
            folder_name = os.path.basename(self.image_folder)
            self.img_path_display.configure(text=folder_name)
            
            # Count images
            image_count = len(self._scan_images_quick())
            if image_count > 0:
                self.img_count_label.configure(text=f"({image_count} images found)")
            else:
                self.img_count_label.configure(text="(no valid images)", text_color="#FF6B6B")
        else:
            self.img_path_display.configure(text="No folder selected")
            self.img_count_label.configure(text="")
        
        # Excel file display
        if self.excel_file:
            file_name = os.path.basename(self.excel_file)
            self.excel_path_display.configure(text=file_name)
            
            # Count rows
            try:
                wb = openpyxl.load_workbook(self.excel_file, read_only=True, data_only=True)
                ws = wb.active
                row_count = ws.max_row - 1  # Exclude header
                headers = [cell.value for cell in ws[1]]
                wb.close()
                
                self.excel_info_label.configure(
                    text=f"({row_count} products, {len(headers)} columns)"
                )
            except Exception as e:
                self.excel_info_label.configure(
                    text=f"(error reading file)",
                    text_color="#FF6B6B"
                )
        else:
            self.excel_path_display.configure(text="No file selected")
            self.excel_info_label.configure(text="")
        
        # Enable/disable Next button
        if self.image_folder and self.excel_file and not self.is_loading:
            self.next_btn.configure(state="normal")
        else:
            self.next_btn.configure(state="disabled")
    
    def _scan_images_quick(self) -> list[str]:
        """Quick scan for image files without loading metadata"""
        if not os.path.isdir(self.image_folder):
            return []
        
        images = []
        for filename in os.listdir(self.image_folder):
            if Path(filename).suffix.lower() in VALID_IMAGE_EXTENSIONS:
                images.append(filename)
        
        return images
    
    # ══════════════════════════════════════════════════════
    # ACTIONS
    # ══════════════════════════════════════════════════════
    
    def on_next(self):
        """Validate and load data, then proceed"""
        
        if not self.image_folder or not self.excel_file:
            messagebox.showerror("Error", "Please select both folder and file")
            return
        
        # Start loading in background thread
        self.is_loading = True
        self.next_btn.configure(state="disabled")
        self.status_label.configure(text="Loading...", text_color="#FFAA00")
        
        thread = threading.Thread(target=self._load_worker, daemon=True)
        thread.start()
    
    def _load_worker(self):
        """Background thread: scan images and read Excel"""
        
        try:
            # Scan images
            self.images = self._scan_images()
            
            if not self.images:
                self.after(0, lambda: messagebox.showerror(
                    "Error", f"No valid images found in:\n{self.image_folder}"
                ))
                self.after(0, self._reset_loading)
                return
            
            # Read Excel
            self.excel_metadata = self._read_excel()
            
            if not self.excel_metadata:
                self.after(0, lambda: messagebox.showerror(
                    "Error", f"Could not read Excel file:\n{self.excel_file}"
                ))
                self.after(0, self._reset_loading)
                return
            
            # Success
            self.after(0, self._on_load_success)
            
        except Exception as e:
            self.after(0, lambda: messagebox.showerror("Error", f"Import failed:\n{str(e)}"))
            self.after(0, self._reset_loading)
    
    def _scan_images(self) -> list[ImageFile]:
        """Scan folder and load image metadata"""
        
        images = []
        
        for filename in sorted(os.listdir(self.image_folder)):
            
            if Path(filename).suffix.lower() not in VALID_IMAGE_EXTENSIONS:
                continue
            
            file_path = os.path.join(self.image_folder, filename)
            
            try:
                file_size = os.path.getsize(file_path)
                image = ImageFile(
                    path=file_path,
                    filename=filename,
                    file_size=file_size
                )
                images.append(image)
            except Exception as e:
                print(f"Warning: Could not load {filename}: {e}")
        
        return images
    
    def _read_excel(self) -> Optional[ExcelMetadata]:
        """Read Excel file and extract metadata"""
        
        try:
            wb = openpyxl.load_workbook(self.excel_file, data_only=False)
            ws = wb.active
            
            # Get headers (first row)
            headers = []
            for cell in ws[1]:
                headers.append(cell.value or "")
            
            # Count data rows
            total_rows = ws.max_row - 1  # Exclude header
            
            wb.close()
            
            metadata = ExcelMetadata(
                file_path=self.excel_file,
                total_rows=total_rows,
                headers=headers
            )
            
            return metadata
            
        except Exception as e:
            print(f"Error reading Excel: {e}")
            return None
    
    def _on_load_success(self):
        """Called when import succeeds"""
        
        self.status_label.configure(
            text=f"✓ Loaded {len(self.images)} images and {self.excel_metadata.total_rows} products",
            text_color="#66BB6A"
        )
        
        # Call completion callback
        if self.on_complete:
            self.on_complete(self.images, self.excel_metadata)
    
    def _reset_loading(self):
        """Reset loading state"""
        
        self.is_loading = False
        self.update_ui()
    
    def on_cancel(self):
        """Cancel and go back"""
        
        if messagebox.askyesno("Cancel", "Discard current import?"):
            self.destroy()
    
    # ══════════════════════════════════════════════════════
    # DATA ACCESS
    # ══════════════════════════════════════════════════════
    
    def get_data(self) -> tuple[list[ImageFile], ExcelMetadata]:
        """Get loaded data"""
        return self.images, self.excel_metadata
