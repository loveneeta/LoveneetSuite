import io, os, re, shutil, subprocess, tempfile, zipfile
from pathlib import Path
from typing import Annotated
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from PIL import Image

app = FastAPI(title="Loveneet Suite Web API")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173","http://127.0.0.1:5173"], allow_methods=["*"], allow_headers=["*"])
EXTS = {".png",".jpg",".jpeg",".webp",".bmp",".tif",".tiff"}

@app.get("/api/health")
def health(): return {"ok": True, "service": "Loveneet Suite"}

def safe_name(name): return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(name).name).strip(" .") or "file"

def zip_folder(folder, filename):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as z:
        for f in Path(folder).rglob("*"):
            if f.is_file(): z.write(f, f.relative_to(folder).as_posix())
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/zip", headers={"Content-Disposition": 'attachment; filename="' + filename + '"'})

@app.post("/api/convert")
async def convert(files: Annotated[list[UploadFile], File()], output_format: Annotated[str, Form()] = "jpg", quality: Annotated[int, Form()] = 90, width: Annotated[int | None, Form()] = None, height: Annotated[int | None, Form()] = None):
    fmt = output_format.lower().lstrip(".")
    if fmt not in {"jpg","jpeg","png","webp","bmp","tiff"}: raise HTTPException(400, "Unsupported format")
    if not 10 <= quality <= 100: raise HTTPException(400, "Quality must be between 10 and 100")
    if (width is None) != (height is None): raise HTTPException(400, "Enter both width and height, or neither")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "converted"; out.mkdir(); used=set(); errors=[]; count=0
        for upload in files:
            name=safe_name(upload.filename or "image")
            if Path(name).suffix.lower() not in EXTS: errors.append(name + ": unsupported format"); continue
            try:
                with Image.open(io.BytesIO(await upload.read())) as source:
                    image=source.copy()
                    if width and height:
                        if width > 12000 or height > 12000 or width < 1 or height < 1: raise ValueError("Invalid dimensions")
                        image=image.resize((width,height), Image.Resampling.LANCZOS)
                    if fmt in {"jpg","jpeg"}:
                        if image.mode in ("RGBA","LA") or (image.mode=="P" and "transparency" in image.info):
                            rgba=image.convert("RGBA"); bg=Image.new("RGB",rgba.size,"white"); bg.paste(rgba,mask=rgba.getchannel("A")); image=bg
                        else: image=image.convert("RGB")
                    ext=".jpg" if fmt in {"jpg","jpeg"} else "."+fmt; stem=Path(name).stem; target_name=stem+ext; n=1
                    while target_name.casefold() in used: target_name=stem+"_"+str(n)+ext; n+=1
                    used.add(target_name.casefold()); opts={"quality":quality} if fmt in {"jpg","jpeg","webp"} else {}
                    image.save(out/target_name, format={"jpg":"JPEG","jpeg":"JPEG","tiff":"TIFF"}.get(fmt,fmt.upper()), **opts); count+=1
            except Exception as exc: errors.append(name + ": " + str(exc))
        if not count: raise HTTPException(400, "No images converted. " + "; ".join(errors[:4]))
        (out/"conversion_report.txt").write_text("Converted: "+str(count)+"\nErrors: "+str(len(errors))+"\n"+"\n".join(errors),encoding="utf-8")
        return zip_folder(out,"converted_images.zip")

@app.post("/api/organize")
async def organize(excel: Annotated[UploadFile, File()], image_archive: Annotated[UploadFile, File()], mode: Annotated[int, Form()] = 2):
    if mode not in {1,2,3}: raise HTTPException(400,"Mode must be 1, 2, or 3")
    if Path(excel.filename or "").suffix.lower() not in {".xlsx",".xlsm"}: raise HTTPException(400,"Upload XLSX or XLSM")
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp); srcdir=root/"images"; out=root/"organized"; srcdir.mkdir(); out.mkdir()
        try:
            with zipfile.ZipFile(io.BytesIO(await image_archive.read())) as z:
                for item in z.infolist():
                    dest=(srcdir/item.filename).resolve()
                    if srcdir.resolve() not in dest.parents and dest != srcdir.resolve(): raise HTTPException(400,"Unsafe ZIP path")
                z.extractall(srcdir)
        except zipfile.BadZipFile: raise HTTPException(400,"Images must be a ZIP archive")
        workbook=root/safe_name(excel.filename or "products.xlsx"); workbook.write_bytes(await excel.read())
        try: wb=load_workbook(workbook,data_only=True,read_only=True)
        except Exception as exc: raise HTTPException(400,"Cannot read Excel: "+str(exc))
        lookup={p.name.casefold():p for p in srcdir.rglob("*") if p.is_file() and p.suffix.lower() in EXTS}
        copied=skipped=0; errors=[]
        def val(ws,row,col): return str(ws.cell(row=row,column=col).value or "").strip()
        def clean(s): return re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",s).strip(" .")
        for ws in wb.worksheets:
            for row in range(2,ws.max_row+1):
                try:
                    cell=ws.cell(row=row,column=5); raw=cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else str(cell.value or "")
                    src=lookup.get(raw.replace("\\","/").split("/")[-1].casefold())
                    if not src: skipped+=1; continue
                    if mode==1: folders=[clean(val(ws,row,c)) for c in range(7,18) if clean(val(ws,row,c))]
                    elif mode==2: folders=["DELETE" if "DELETE" in val(ws,row,7).upper() else "KEEP",clean(val(ws,row,17)) or "SOLE"]
                    else: folders=["DELETE" if "DELETE" in val(ws,row,7).upper() else "KEEP",clean(val(ws,row,8)) or "UNCATEGORIZED"]
                    target_dir=out.joinpath(*folders); target_dir.mkdir(parents=True,exist_ok=True)
                    parts=[clean(val(ws,row,c)).replace(" ","-") for c in [3,10,11,13,14,12,15,16,26] if clean(val(ws,row,c))]
                    stem="_".join(parts) or "unnamed_"+str(row); target=target_dir/(stem+src.suffix); n=1
                    while target.exists(): target=target_dir/(stem+"_"+str(n)+src.suffix); n+=1
                    shutil.copy2(src,target); copied+=1
                except Exception as exc: errors.append(ws.title+" row "+str(row)+": "+str(exc))
        wb.close(); (out/"organizer_report.txt").write_text("Copied: "+str(copied)+"\nSkipped: "+str(skipped)+"\nErrors: "+str(len(errors))+"\n"+"\n".join(errors),encoding="utf-8")
        return zip_folder(out,"organized_images.zip")

@app.post("/api/ppt-export")
async def ppt_export(files: Annotated[list[UploadFile], File()], export_jpg: Annotated[bool, Form()] = True, export_png: Annotated[bool, Form()] = False, export_webp: Annotated[bool, Form()] = True, quality: Annotated[int, Form()] = 95):
    if not any([export_jpg,export_png,export_webp]): raise HTTPException(400,"Select an output format")
    if shutil.which("libreoffice") is None and shutil.which("soffice") is None: raise HTTPException(501,"PowerPoint export needs LibreOffice installed on the server")
    import fitz
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp); out=root/"exports"; out.mkdir()
        for upload in files:
            name=safe_name(upload.filename or "slides.pptx")
            if Path(name).suffix.lower() not in {".ppt",".pptx"}: continue
            ppt=root/name; ppt.write_bytes(await upload.read()); result=subprocess.run([shutil.which("libreoffice") or shutil.which("soffice"),"--headless","--convert-to","pdf","--outdir",str(root),str(ppt)],capture_output=True,text=True,timeout=180)
            pdf=root/(ppt.stem+".pdf")
            if result.returncode or not pdf.exists(): raise HTTPException(500,"PowerPoint conversion failed: "+result.stderr[-300:])
            folder=out/ppt.stem; folder.mkdir(); prs=Presentation(str(ppt))
            with fitz.open(pdf) as doc:
                for i,page in enumerate(doc):
                    label="Slide_"+str(i+1).zfill(3)
                    if i<len(prs.slides):
                        found=[]
                        for shape in prs.slides[i].shapes:
                            if getattr(shape,"has_text_frame",False) and "|" in shape.text: found.append((shape.top,shape.left,shape.text.split("|",1)[0].strip()))
                        if found: label=safe_name(sorted(found)[0][2]) or label
                    pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False); image=Image.open(io.BytesIO(pix.tobytes("png")))
                    if export_png: image.save(folder/(label+".png"),"PNG")
                    if export_jpg: image.convert("RGB").save(folder/(label+".jpg"),"JPEG",quality=95)
                    if export_webp: image.save(folder/(label+".webp"),"WEBP",quality=quality)
            pdf.unlink(missing_ok=True)
        return zip_folder(out,"powerpoint_exports.zip")
