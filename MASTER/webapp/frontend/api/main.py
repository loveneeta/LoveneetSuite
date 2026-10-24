import io, json, os, re, shutil, subprocess, tempfile, zipfile
from pathlib import Path
from typing import Annotated
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from PIL import Image

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

app = FastAPI(title="Loveneet Suite Web API")

allowed_origins = [
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    "http://localhost:5174",
    "http://127.0.0.1:5174",
    "http://localhost:3000",
    "http://127.0.0.1:3000",
]
env_origins = os.getenv("CORS_ORIGINS")
if env_origins:
    allowed_origins.extend([o.strip() for o in env_origins.split(",") if o.strip()])

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins if os.getenv("ALLOW_ALL_CORS", "true").lower() != "true" else ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
EXTS = {".png",".jpg",".jpeg",".webp",".bmp",".tif",".tiff"}

@app.get("/api/health")
def health(): return {"ok": True, "service": "Loveneet Suite"}

def safe_name(name): return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", Path(name).name).strip(" .") or "file"

def zip_folder(folder, filename):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as z:
        for f in Path(folder).rglob("*"):
            if f.is_file(): z.write(f, f.relative_to(folder).as_posix())
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/zip", headers={"Content-Disposition": 'attachment; filename="' + filename + '"'})

@app.post("/api/convert")
async def convert(files: Annotated[list[UploadFile], File()], output_format: Annotated[str, Form()] = "jpg", quality: Annotated[int, Form()] = 90, width: Annotated[int | None, Form()] = None, height: Annotated[int | None, Form()] = None):
    fmt = output_format.lower().lstrip(".")
    if fmt not in {"jpg","jpeg","png","webp","bmp","tiff"}: raise HTTPException(400, "Unsupported format")
    if not 10 <= quality <= 100: raise HTTPException(400, "Quality must be between 10 and 100")
    if (width is None) != (height is None): raise HTTPException(400, "Enter both width and height, or neither")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "converted"; out.mkdir(); used=set(); errors=[]; count=0
        for upload in files:
            name=safe_name(upload.filename or "image")
            if Path(name).suffix.lower() not in EXTS: errors.append(name + ": unsupported format"); continue
            try:
                with Image.open(io.BytesIO(await upload.read())) as source:
                    image=source.copy()
                    if width and height:
                        if width > 12000 or height > 12000 or width < 1 or height < 1: raise ValueError("Invalid dimensions")
                        image=image.resize((width,height), Image.Resampling.LANCZOS)
                    if fmt in {"jpg","jpeg"}:
                        if image.mode in ("RGBA","LA") or (image.mode=="P" and "transparency" in image.info):
                            rgba=image.convert("RGBA"); bg=Image.new("RGB",rgba.size,"white"); bg.paste(rgba,mask=rgba.getchannel("A")); image=bg
                        else: image=image.convert("RGB")
                    ext=".jpg" if fmt in {"jpg","jpeg"} else "."+fmt; stem=Path(name).stem; target_name=stem+ext; n=1
                    while target_name.casefold() in used: target_name=stem+"_"+str(n)+ext; n+=1
                    used.add(target_name.casefold()); opts={"quality":quality} if fmt in {"jpg","jpeg","webp"} else {}
                    image.save(out/target_name, format={"jpg":"JPEG","jpeg":"JPEG","tiff":"TIFF"}.get(fmt,fmt.upper()), **opts); count+=1
            except Exception as exc: errors.append(name + ": " + str(exc))
        if not count: raise HTTPException(400, "No images converted. " + "; ".join(errors[:4]))
        (out/"conversion_report.txt").write_text("Converted: "+str(count)+"\nErrors: "+str(len(errors))+"\n"+"\n".join(errors),encoding="utf-8")
        return zip_folder(out,"converted_images.zip")

@app.post("/api/organize")
async def organize(excel: Annotated[UploadFile, File()], image_archive: Annotated[UploadFile, File()], mode: Annotated[int, Form()] = 2):
    if mode not in {1,2,3}: raise HTTPException(400,"Mode must be 1, 2, or 3")
    if Path(excel.filename or "").suffix.lower() not in {".xlsx",".xlsm"}: raise HTTPException(400,"Upload XLSX or XLSM")
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp); srcdir=root/"images"; out=root/"organized"; srcdir.mkdir(); out.mkdir()
        try:
            with zipfile.ZipFile(io.BytesIO(await image_archive.read())) as z:
                for item in z.infolist():
                    dest=(srcdir/item.filename).resolve()
                    if srcdir.resolve() not in dest.parents and dest != srcdir.resolve(): raise HTTPException(400,"Unsafe ZIP path")
                z.extractall(srcdir)
        except zipfile.BadZipFile: raise HTTPException(400,"Images must be a ZIP archive")
        workbook=root/safe_name(excel.filename or "products.xlsx"); workbook.write_bytes(await excel.read())
        try: wb=load_workbook(workbook,data_only=True,read_only=True)
        except Exception as exc: raise HTTPException(400,"Cannot read Excel: "+str(exc))
        lookup={p.name.casefold():p for p in srcdir.rglob("*") if p.is_file() and p.suffix.lower() in EXTS}
        copied=skipped=0; errors=[]
        def val(ws,row,col): return str(ws.cell(row=row,column=col).value or "").strip()
        def clean(s): return re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",s).strip(" .")
        for ws in wb.worksheets:
            for row in range(2,ws.max_row+1):
                try:
                    cell=ws.cell(row=row,column=5); raw=cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else str(cell.value or "")
                    src=lookup.get(raw.replace("\\","/").split("/")[-1].casefold())
                    if not src: skipped+=1; continue
                    if mode==1: folders=[clean(val(ws,row,c)) for c in range(7,18) if clean(val(ws,row,c))]
                    elif mode==2: folders=["DELETE" if "DELETE" in val(ws,row,7).upper() else "KEEP",clean(val(ws,row,17)) or "SOLE"]
                    else: folders=["DELETE" if "DELETE" in val(ws,row,7).upper() else "KEEP",clean(val(ws,row,8)) or "UNCATEGORIZED"]
                    target_dir=out.joinpath(*folders); target_dir.mkdir(parents=True,exist_ok=True)
                    parts=[clean(val(ws,row,c)).replace(" ","-") for c in [3,10,11,13,14,12,15,16,26] if clean(val(ws,row,c))]
                    stem="_".join(parts) or "unnamed_"+str(row); target=target_dir/(stem+src.suffix); n=1
                    while target.exists(): target=target_dir/(stem+"_"+str(n)+src.suffix); n+=1
                    shutil.copy2(src,target); copied+=1
                except Exception as exc: errors.append(ws.title+" row "+str(row)+": "+str(exc))
        wb.close(); (out/"organizer_report.txt").write_text("Copied: "+str(copied)+"\nSkipped: "+str(skipped)+"\nErrors: "+str(len(errors))+"\n"+"\n".join(errors),encoding="utf-8")
        return zip_folder(out,"organized_images.zip")

@app.post("/api/ppt-export")
async def ppt_export(files: Annotated[list[UploadFile], File()], export_jpg: Annotated[bool, Form()] = True, export_png: Annotated[bool, Form()] = False, export_webp: Annotated[bool, Form()] = True, quality: Annotated[int, Form()] = 95):
    if not any([export_jpg,export_png,export_webp]): raise HTTPException(400,"Select an output format")
    if shutil.which("libreoffice") is None and shutil.which("soffice") is None: raise HTTPException(501,"PowerPoint export needs LibreOffice installed on the server")
    if fitz is None: raise HTTPException(501,"PDF slide rendering requires PyMuPDF (pymupdf). Install pymupdf on the server.")
    from pptx import Presentation
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp); out=root/"exports"; out.mkdir()
        for upload in files:
            name=safe_name(upload.filename or "slides.pptx")
            if Path(name).suffix.lower() not in {".ppt",".pptx"}: continue
            ppt=root/name; ppt.write_bytes(await upload.read()); result=subprocess.run([shutil.which("libreoffice") or shutil.which("soffice"),"--headless","--convert-to","pdf","--outdir",str(root),str(ppt)],capture_output=True,text=True,timeout=180)
            pdf=root/(ppt.stem+".pdf")
            if result.returncode or not pdf.exists(): raise HTTPException(500,"PowerPoint conversion failed: "+result.stderr[-300:])
            folder=out/ppt.stem; folder.mkdir(); prs=Presentation(str(ppt))
            with fitz.open(pdf) as doc:
                for i,page in enumerate(doc):
                    label="Slide_"+str(i+1).zfill(3)
                    if i<len(prs.slides):
                        found=[]
                        for shape in prs.slides[i].shapes:
                            if getattr(shape,"has_text_frame",False) and "|" in shape.text: found.append((shape.top,shape.left,shape.text.split("|",1)[0].strip()))
                        if found: label=safe_name(sorted(found)[0][2]) or label
                    pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False); image=Image.open(io.BytesIO(pix.tobytes("png")))
                    if export_png: image.save(folder/(label+".png"),"PNG")
                    if export_jpg: image.convert("RGB").save(folder/(label+".jpg"),"JPEG",quality=95)
                    if export_webp: image.save(folder/(label+".webp"),"WEBP",quality=quality)
            pdf.unlink(missing_ok=True)
        return zip_folder(out,"powerpoint_exports.zip")


@app.post("/api/excel-to-ppt")
async def excel_to_ppt(
    excel: Annotated[UploadFile, File()],
    mapping: Annotated[str, Form()],
    layout: Annotated[str, Form()],
    style: Annotated[str, Form()],
    logo: UploadFile | None = File(default=None),
):
    """Build a product catalogue presentation from an XLSX workbook and its embedded cell images."""
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    if Path(excel.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(400, "Upload an .xlsx workbook.")
    try:
        mapping_data = json.loads(mapping)
        slots = json.loads(layout)
        style_data = json.loads(style)
    except (ValueError, TypeError):
        raise HTTPException(400, "Invalid mapping, layout, or style settings.")
    if not isinstance(slots, list) or not 1 <= len(slots) <= 6:
        raise HTTPException(400, "Choose between 1 and 6 product slots per slide.")
    try:
        for slot in slots:
            for rect_name in ("img", "text"):
                rect = slot[rect_name]
                if len(rect) != 4:
                    raise ValueError()
                x1, y1, x2, y2 = [float(v) for v in rect]
                if not (0 <= x1 < x2 <= 800 and 0 <= y1 < y2 <= 450):
                    raise ValueError()
        font_size = max(8, min(48, int(style_data.get("font_size", 18))))
        font_name = str(style_data.get("font_name", "Arial"))[:80]
        font_color = str(style_data.get("font_color", "000000")).replace("#", "").upper()
        background_color = str(style_data.get("background_color", "FFFFFF")).replace("#", "").upper()
        aspect_ratio = str(style_data.get("aspect_ratio", "16:9"))
        ratio_map = {"16:9": 16 / 9, "4:3": 4 / 3, "A4 Landscape": 297 / 210, "A4 Portrait": 210 / 297}
        if aspect_ratio not in ratio_map:
            raise ValueError()
        if not re.fullmatch(r"[0-9A-F]{6}", font_color) or not re.fullmatch(r"[0-9A-F]{6}", background_color):
            raise ValueError()
    except (KeyError, TypeError, ValueError):
        raise HTTPException(400, "Invalid layout coordinates or style settings. Coordinates must fit within 800 × 450.")

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        workbook_path = root / safe_name(excel.filename or "products.xlsx")
        workbook_path.write_bytes(await excel.read())
        try:
            wb = load_workbook(workbook_path, data_only=True)
        except Exception as exc:
            raise HTTPException(400, "Cannot read Excel workbook: " + str(exc))
        if not wb.worksheets:
            raise HTTPException(400, "Workbook contains no worksheets.")
        ws = wb.active
        headers = {}
        for cell in ws[1]:
            if cell.value is not None:
                headers[str(cell.value).strip().casefold()] = cell.column

        def mapped_col(role, *fallbacks):
            selected = str(mapping_data.get(role) or "").strip().casefold()
            if selected and selected in headers:
                return headers[selected]
            for fallback in fallbacks:
                if fallback in headers:
                    return headers[fallback]
            return None

        code_col = mapped_col("code", "sku", "item code", "product code", "code")
        desc_col = mapped_col("description", "description", "product name", "product", "desc")
        price_col = mapped_col("price", "price", "mrp", "rate")
        length_col = mapped_col("length", "length", "l")
        width_col = mapped_col("width", "width", "breadth", "w")
        height_col = mapped_col("height", "height", "h")
        image_col = mapped_col("image", "image", "photo", "picture", "img")
        if not code_col:
            wb.close()
            raise HTTPException(400, "Map a product code / SKU column.")
        
        # openpyxl keeps worksheet pictures anchored to zero-based row/column coordinates.
        image_by_cell = {}
        try:
            for picture in getattr(ws, "_images", []):
                anchor = getattr(picture, "anchor", None)
                marker = getattr(anchor, "_from", None)
                if marker is None:
                    continue
                row_number, column_number = marker.row + 1, marker.col + 1
                if image_col is None or column_number == image_col:
                    image_by_cell.setdefault(row_number, picture)
        except Exception:
            image_by_cell = {}

        def cell_value(row_number, column_number):
            if not column_number:
                return ""
            value = ws.cell(row=row_number, column=column_number).value
            return "" if value is None else str(value).strip()

        products = []
        for row_number in range(2, ws.max_row + 1):
            code = cell_value(row_number, code_col)
            if not code:
                continue
            products.append({
                "row": row_number,
                "code": code,
                "description": cell_value(row_number, desc_col),
                "price": cell_value(row_number, price_col),
                "length": cell_value(row_number, length_col),
                "width": cell_value(row_number, width_col),
                "height": cell_value(row_number, height_col),
                "picture": image_by_cell.get(row_number),
            })
        if not products:
            wb.close()
            raise HTTPException(400, "No product rows found. Check the product code column mapping.")
        if len(products) > 5000:
            wb.close()
            raise HTTPException(400, "This workbook has more than 5,000 products. Split it into smaller workbooks.")

        prs = Presentation()
        slide_width_in = 8.0
        slide_height_in = slide_width_in / ratio_map[aspect_ratio]
        prs.slide_width = Inches(slide_width_in)
        prs.slide_height = Inches(slide_height_in)
        blank_layout = prs.slide_layouts[6]
        header_text = str(style_data.get("header_text", ""))[:180]
        show_header = bool(style_data.get("show_header", True))
        show_rule = bool(style_data.get("show_rule", True))
        logo_bytes = None
        if logo is not None and logo.filename:
            logo_ext = Path(logo.filename).suffix.lower()
            if logo_ext not in {".png", ".jpg", ".jpeg", ".webp"}:
                wb.close()
                raise HTTPException(400, "Logo must be PNG, JPG, JPEG, or WEBP.")
            logo_bytes = await logo.read()
            if len(logo_bytes) > 5 * 1024 * 1024:
                wb.close()
                raise HTTPException(400, "Logo must be 5 MB or smaller.")
        auto_center_image = bool(style_data.get("auto_center_image", True))
        image_fit = str(style_data.get("image_fit", "contain"))
        if image_fit not in {"contain", "cover"}:
            image_fit = "contain"
        for product_index, product in enumerate(products):
            slot_index = product_index % len(slots)
            if slot_index == 0:
                slide = prs.slides.add_slide(blank_layout)
                fill = slide.background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor.from_string(background_color)
                if show_header and header_text:
                    head = slide.shapes.add_textbox(Inches(0.28), Inches(0.10), Inches(slide_width_in - 0.56), Inches(0.38))
                    tf = head.text_frame
                    tf.clear()
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.LEFT
                    run = p.add_run()
                    run.text = header_text
                    run.font.name = font_name
                    run.font.size = Pt(max(10, font_size))
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(font_color)
                if show_rule:
                    line_y = 0.58 if show_header else 0.22
                    rule = slide.shapes.add_shape(1, Inches(0.28), Inches(line_y), Inches(slide_width_in - 0.56), Inches(0.012))
                    rule.fill.solid()
                    rule.fill.fore_color.rgb = RGBColor.from_string(font_color)
                    rule.line.fill.background()
                if logo_bytes:
                    try:
                        slide.shapes.add_picture(io.BytesIO(logo_bytes), Inches(slide_width_in - 1.05), Inches(0.08), width=Inches(0.72), height=Inches(0.42))
                    except Exception:
                        pass
            slot = slots[slot_index]
            if product["picture"] is not None:
                try:
                    img_bytes = product["picture"]._data()
                    from PIL import Image, ImageOps
                    image = Image.open(io.BytesIO(img_bytes)).convert("RGBA")
                    box_w = max(1, int(float(slot["img"][2]) - float(slot["img"][0])))
                    box_h = max(1, int(float(slot["img"][3]) - float(slot["img"][1])))
                    if image_fit == "cover":
                        fitted = ImageOps.fit(image, (box_w * 2, box_h * 2), method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))
                    else:
                        fitted = Image.new("RGBA", (box_w * 2, box_h * 2), (255,255,255,0))
                        contained = ImageOps.contain(image, (box_w * 2, box_h * 2), method=Image.Resampling.LANCZOS)
                        fitted.alpha_composite(contained, ((fitted.width-contained.width)//2, (fitted.height-contained.height)//2))
                    image_buffer = io.BytesIO()
                    fitted.save(image_buffer, format="PNG")
                    image_buffer.seek(0)
                    img_bytes = image_buffer.getvalue()
                    x1, y1, x2, y2 = [float(v) for v in slot["img"]]
                    tx1, _, tx2, _ = [float(v) for v in slot["text"]]
                    if auto_center_image:
                        image_width = x2 - x1
                        image_center = (tx1 + tx2) / 2
                        x1 = max(0, min(800 - image_width, image_center - image_width / 2))
                        x2 = x1 + image_width
                    x1_in, y1_in = (x1 / 800) * slide_width_in, (y1 / 450) * slide_height_in
                    width_in, height_in = ((x2 - x1) / 800) * slide_width_in, ((y2 - y1) / 450) * slide_height_in
                    slide.shapes.add_picture(io.BytesIO(img_bytes), Inches(x1_in), Inches(y1_in), width=Inches(width_in), height=Inches(height_in))
                except Exception:
                    pass

            tx1, ty1, tx2, ty2 = [float(v) for v in slot["text"]]
            tx1_in, ty1_in = (tx1 / 800) * slide_width_in, (ty1 / 450) * slide_height_in
            text_width_in, text_height_in = ((tx2 - tx1) / 800) * slide_width_in, ((ty2 - ty1) / 450) * slide_height_in
            textbox = slide.shapes.add_textbox(Inches(tx1_in), Inches(ty1_in), Inches(text_width_in), Inches(text_height_in))
            text_frame = textbox.text_frame
            text_frame.clear()
            text_frame.word_wrap = True
            line1 = product["code"] + (" | " + product["description"] if product["description"] else "")
            dimensions = ""
            if product["length"] and product["width"] and product["height"]:
                dimensions = f'{product["length"]} x {product["width"]} x {product["height"]}'
                try:
                    dims_in = [round(float(v) / 2.54, 1) for v in (product["length"], product["width"], product["height"])]
                    dimensions += " CM | " + " x ".join(str(v) for v in dims_in) + " IN"
                except (TypeError, ValueError):
                    dimensions += " CM"
            lines = [line1]
            if dimensions:
                lines.append(dimensions)
            if product["price"]:
                lines.append(product["price"])
            for line_index, line in enumerate(lines):
                paragraph = text_frame.paragraphs[0] if line_index == 0 else text_frame.add_paragraph()
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = line
                run.font.name = font_name
                run.font.size = Pt(font_size if line_index == 0 else (font_size - 2 if line_index == 1 else font_size - 1))
                run.font.bold = line_index in (0, len(lines)-1) and (line_index == 0 or len(lines) > 1)
                run.font.color.rgb = RGBColor.from_string(font_color)

        wb.close()
        output = root / "product_catalogue.pptx"
        prs.save(output)
        from fastapi.responses import FileResponse
        return StreamingResponse(io.BytesIO(output.read_bytes()), media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", headers={"Content-Disposition": 'attachment; filename="product_catalogue.pptx"'})


@app.post("/api/ppt-excel/extract")
async def ppt_excel_extract(file: UploadFile = File(...)):
    import base64

    from pptx import Presentation
    ext = Path(file.filename or "").suffix.lower()
    if ext not in {".pptx", ".ppt"}:
        raise HTTPException(400, "Upload a .pptx or .ppt PowerPoint file.")
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        original = root / safe_name(file.filename or "presentation.pptx")
        original.write_bytes(await file.read())
        pptx_path = original
        converter = shutil.which("libreoffice") or shutil.which("soffice")
        if ext == ".ppt":
            if not converter:
                raise HTTPException(501, "Legacy .ppt conversion requires LibreOffice on the server. Save the file as .pptx and retry.")
            result = subprocess.run([converter, "--headless", "--convert-to", "pptx", "--outdir", str(root), str(original)], capture_output=True, text=True, timeout=180)
            pptx_path = root / (original.stem + ".pptx")
            if result.returncode or not pptx_path.exists():
                raise HTTPException(422, "Could not convert legacy PPT file: " + (result.stderr or result.stdout)[-300:])
        try:
            prs = Presentation(str(pptx_path))
        except Exception as exc:
            raise HTTPException(422, "PowerPoint could not be opened: " + str(exc)[:200])
        slides = []
        for si, slide in enumerate(prs.slides, start=1):
            texts, tables, images, title = [], [], [], ""
            for ix, shape in enumerate(slide.shapes):
                try:
                    if getattr(shape, "has_text_frame", False):
                        val = shape.text.strip()
                        if val:
                            texts.append(val)
                            if getattr(shape, "is_placeholder", False) and not title:
                                title = val
                    if getattr(shape, "has_table", False):
                        table = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                        tables.append(table)
                        texts.extend([" | ".join(row) for row in table if any(row)])
                    if getattr(shape, "shape_type", None) == 13:
                        pic = shape.image
                        sw = float(prs.slide_width) if getattr(prs, "slide_width", None) else 9144000.0
                        sh = float(prs.slide_height) if getattr(prs, "slide_height", None) else 5143500.0
                        rx = float(shape.left) / sw
                        ry = float(shape.top) / sh
                        rw = float(shape.width) / sw
                        rh = float(shape.height) / sh
                        is_top_corner = ry < 0.22 and (rx > 0.65 or rx < 0.28) and rw < 0.35 and rh < 0.25
                        is_footer = ry > 0.80 or rh < 0.025 or (rw > 0.4 and rh < 0.035)
                        is_logo = is_top_corner or is_footer
                        images.append({"id": "s" + str(si) + "i" + str(ix), "name": "slide_" + str(si) + "_image_" + str(len(images)+1) + "." + (pic.ext or "png"), "mime": pic.content_type or "image/png", "data": base64.b64encode(pic.blob).decode("ascii"), "isLogo": is_logo, "role": "logo" if is_logo else "product"})
                except Exception:
                    continue
            preview = None
            slides.append({"slideNumber":si,"sourceFile":file.filename or "presentation.pptx","title":title or (texts[0][:100] if texts else "Slide "+str(si)),"text":texts,"tables":tables,"images":images,"preview":preview})
        if not slides:
            raise HTTPException(422, "No slides were found.")
        return {"sourceFile":file.filename,"slideCount":len(slides),"slides":slides}


@app.post("/api/ppt-excel/export")
async def ppt_excel_export(payload: dict):
    import base64
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    if not isinstance(payload, dict) or not isinstance(payload.get("columns"), list) or not isinstance(payload.get("rows"), list):
        raise HTTPException(400, "Invalid workbook payload.")
    columns, rows = payload["columns"], payload["rows"]
    if not columns or len(columns) > 100:
        raise HTTPException(400, "Configure between 1 and 100 columns.")
    if len(rows) > 1500:
        raise HTTPException(413, "Export up to 1,500 rows at a time. Split the presentation into smaller batches.")
    headers = [str(c.get("header") or "Column")[:200] for c in columns]
    if payload.get("include_source"):
        headers += ["Source File", "Slide Number"]
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp); out = root / "PPT_to_Excel_with_Images.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "PPT Extract"
        ws.append(headers); ws.freeze_panes = "A2"; ws.auto_filter.ref = "A1:" + get_column_letter(len(headers)) + str(len(rows)+1)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="263D70"); cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        for ci, col in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(ci)].width = 22 if col.get("type") == "image" else min(42,max(15,len(headers[ci-1])+3))
        if payload.get("include_source"):
            ws.column_dimensions[get_column_letter(len(headers)-1)].width = 30
        errors = []
        for ri, row in enumerate(rows, 2):
            vals, pics = row.get("values") or {}, row.get("images") or {}
            for ci, col in enumerate(columns, 1):
                if col.get("type") == "image" and payload.get("include_images", True):
                    ws.cell(row=ri,column=ci,value="")
                    picdata = pics.get(col.get("id"))
                    if picdata and picdata.get("data"):
                        try:
                            raw = base64.b64decode(picdata["data"], validate=True)
                            imgpath = root / ("img_" + str(ri) + "_" + str(ci) + ".dat"); imgpath.write_bytes(raw)
                            pic = XLImage(str(imgpath)); scale = min(128/max(pic.width,1),88/max(pic.height,1),1)
                            pic.width=max(1,int(pic.width*scale)); pic.height=max(1,int(pic.height*scale)); pic.anchor=ws.cell(row=ri,column=ci).coordinate; ws.add_image(pic)
                            ws.row_dimensions[ri].height=max(ws.row_dimensions[ri].height or 15,70)
                        except Exception as exc:
                            errors.append("Row "+str(ri)+", "+str(col.get("header"))+": "+str(exc)[:150])
                else:
                    ws.cell(row=ri,column=ci,value=str(vals.get(col.get("id")) or "")[:32000])
            if payload.get("include_source"):
                ws.cell(row=ri,column=len(columns)+1,value=str(row.get("source_file") or "")[:500])
                ws.cell(row=ri,column=len(columns)+2,value=row.get("slide_number") or "")
        if errors:
            ew = wb.create_sheet("Processing Errors"); ew.append(["Error"])
            for e in errors: ew.append([e])
            ew.column_dimensions["A"].width=100
        wb.save(out)
        try:
            check = load_workbook(out)
            check.close()
        except Exception as exc:
            raise HTTPException(500, "Generated workbook validation failed: " + str(exc)[:180])
        return StreamingResponse(io.BytesIO(out.read_bytes()), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="PPT_to_Excel_with_Images.xlsx"'})
